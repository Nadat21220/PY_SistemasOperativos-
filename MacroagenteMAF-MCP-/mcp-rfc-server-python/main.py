#!/usr/bin/env python3
"""
MCP Server RFC Management System - Python
Migrado de TypeScript a Python para mayor compatibilidad y simplicidad
"""

import logging
import json
import sys
import os
from typing import Optional

import openpyxl

from mcp.server.fastmcp import FastMCP

# Importar configuración y utilidades
from config.settings import (
    MCP_SERVER_NAME,
    MCP_SERVER_VERSION,
    RFC_FILES_DIR,
    validate_config
)
from utils.database import (
    get_db_pool,
    query_select,
    query_insert,
    query_insert_returning,
    log_audit,
    close_db_pool
)
from utils.sql_guard import validate_query, SQLGuardError

# ===== CONFIGURAR LOGGING =====
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# ===== INICIAR MCP SERVER =====
mcp = FastMCP(MCP_SERVER_NAME)


def _read_xlsx_to_markdown(file_path: str) -> str:
    """Lee un .xlsx y devuelve su contenido real (todas las hojas) en formato markdown."""
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"## Hoja: {sheet_name}\n")
        for row in ws.iter_rows():
            values = [str(cell.value).strip() for cell in row if cell.value is not None and str(cell.value).strip()]
            if values:
                parts.append(" | ".join(values))
        parts.append("")
    wb.close()
    return "\n".join(parts)

# ===== HERRAMIENTA 1: Control PostgreSQL =====
@mcp.tool()
def tool_postgres_control(action: str = "", query: str = "", params: Optional[list] = None) -> str:
    """
    Ejecuta SQL en PostgreSQL. Usa SOLO estas tablas/columnas (no existen "rfc", "status" ni otras):

    rfc_records: project_id* UNIQUE, environment[Dev|QA|PROD], project_name, impacted_platforms
    directorio_personal: empleado_id* uuid, nombre, apellido, departamento, email
    rfc_responsables: asignacion_id*, project_id->rfc_records, empleado_id->directorio_personal, rol_asignado
    rfc_document_index: index_id*, file_name, file_type, project_id->rfc_records
    rfc_markdown_index: markdown_id*, rfc_file_id->rfc_document_index, project_id, markdown_content, conversion_status, sync_status, created_by
    query_audit: audit_id*, query_type, executed_query, executed_at
    rfc_audit_logs: audit_log_id*, project_id, action_type, performed_by, details_json, performed_at

    Para nombres de responsables usa JOIN (no dos queries separadas):
        SELECT dp.nombre, dp.apellido, rr.rol_asignado FROM rfc_responsables rr
        JOIN directorio_personal dp ON dp.empleado_id = rr.empleado_id WHERE rr.project_id = 'td189-bf25'

    Args:
        action: "select" para leer, "insert" para INSERT/UPDATE
        query: SQL usando SOLO las tablas de arriba
        params: lista de valores para placeholders %s

    Returns:
        JSON con resultados o error.
    """
    if not action or not query:
        return json.dumps({
            "success": False,
            "error": 'action y query son obligatorios. Ejemplo: {"action":"select","query":"SELECT ..."}'
        })

    try:
        validate_query(action, query)
    except SQLGuardError as e:
        return json.dumps({"success": False, "error": str(e)})

    try:
        if action == "select":
            results = query_select(query, params or [])
            return json.dumps(results, indent=2, default=str)
        elif action == "insert":
            rowcount = query_insert(query, params or [])
            return f"✓ Operación exitosa - {rowcount} filas afectadas"
        else:
            return f"✗ Acción no reconocida: {action}"
    except Exception as e:
        logger.error(f"Error PostgreSQL: {e}")
        return f"✗ Error BD: {str(e)}"

# ===== HERRAMIENTA 2: Convertir RFC a Markdown =====
@mcp.tool()
def tool_convert_rfc_to_markdown(
    project_id: str,
    rfc_file_id: str,
    requested_by: str
) -> str:
    """
    Convertir documento RFC (Excel) a Markdown para mejor procesamiento del LLM.

    Args: project_id (ej td189-bf25), rfc_file_id (uuid en rfc_document_index), requested_by (email)
    """
    try:
        # Verificar que el archivo este indexado para este proyecto
        doc = query_select(
            "SELECT index_id, file_name, file_type FROM rfc_document_index WHERE index_id = %s AND project_id = %s",
            [rfc_file_id, project_id]
        )
        if not doc:
            return json.dumps({
                "success": False,
                "error": "Archivo no encontrado en rfc_document_index para este proyecto"
            })
        doc = doc[0]

        rfc = query_select(
            "SELECT project_name, environment FROM rfc_records WHERE project_id = %s",
            [project_id]
        )
        if not rfc:
            return json.dumps({"success": False, "error": "RFC no encontrado"})
        rfc = rfc[0]

        markdown_file_name = doc["file_name"].rsplit(".", 1)[0] + ".md"

        file_path = os.path.join(RFC_FILES_DIR, doc["file_name"])
        header = (
            f"# RFC: {rfc['project_name']}\n\n"
            f"- Proyecto: {project_id}\n"
            f"- Entorno: {rfc['environment']}\n"
            f"- Archivo origen: {doc['file_name']}\n\n"
        )

        if not os.path.isfile(file_path):
            return json.dumps({
                "success": False,
                "error": f"Archivo fisico no encontrado en {file_path}. Verifica que este montado en el volumen RFC."
            })

        if doc["file_type"].lower() == "xlsx":
            body = _read_xlsx_to_markdown(file_path)
        else:
            return json.dumps({
                "success": False,
                "error": f"Tipo de archivo '{doc['file_type']}' aun no soportado para conversion (solo xlsx por ahora)"
            })

        markdown_text = header + body
        content_bytes = markdown_text.encode("utf-8")

        result = query_insert_returning(
            """
            INSERT INTO rfc_markdown_index
                (rfc_file_id, project_id, original_file_name, markdown_file_name,
                 markdown_content, conversion_status, created_by, conversion_completed_at,
                 markdown_size_bytes, markdown_lines)
            VALUES (%s, %s, %s, %s, %s, 'converted', %s, now(), %s, %s)
            RETURNING markdown_id
            """,
            [
                rfc_file_id, project_id, doc["file_name"], markdown_file_name,
                content_bytes, requested_by, len(content_bytes), markdown_text.count("\n") + 1
            ]
        )

        log_audit(
            project_id,
            "RFC_CONVERT_COMPLETED",
            requested_by,
            {"rfc_file_id": rfc_file_id, "markdown_id": str(result["markdown_id"])}
        )

        return json.dumps({
            "success": True,
            "markdown_id": str(result["markdown_id"]),
            "status": "converted",
            "project_id": project_id
        })
    except Exception as e:
        logger.error(f"Error conversión RFC: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== HERRAMIENTA 3: Obtener Markdown Convertido =====
@mcp.tool()
def tool_get_markdown(
    project_id: str,
    requested_by: str
) -> str:
    """
    Obtener documento Markdown convertido de un RFC.
    Args: project_id (ej td189-bf25), requested_by (email)
    """
    try:
        sql = """
        SELECT
            original_file_name,
            markdown_content,
            conversion_status,
            created_at
        FROM rfc_markdown_index
        WHERE project_id = %s
        ORDER BY created_at DESC
        LIMIT 1
        """

        result = query_select(sql, [project_id])

        if not result:
            return json.dumps({
                "success": False,
                "error": "No se encontró Markdown convertido para este proyecto"
            })

        markdown = result[0]

        log_audit(project_id, "MARKDOWN_RETRIEVED", requested_by, {})

        return json.dumps({
            "success": True,
            "original_file_name": markdown.get("original_file_name"),
            "markdown_content": markdown.get("markdown_content"),
            "conversion_status": markdown.get("conversion_status"),
            "created_at": str(markdown.get("created_at"))
        })
    except Exception as e:
        logger.error(f"Error obteniendo Markdown: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== HERRAMIENTA 4: SharePoint - Crear Documento =====
@mcp.tool()
def tool_sharepoint_create(
    markdown_id: str,
    target_folder_id: Optional[str] = None,
    document_name: Optional[str] = None,
    requested_by: Optional[str] = None
) -> str:
    """
    Registra la creacion de un documento en el indice de gestion documental
    (SharePoint no conectado aun: se guarda en PostgreSQL).
    Args: markdown_id (uuid), target_folder_id, document_name, requested_by (email)
    """
    try:
        existing = query_select(
            "SELECT markdown_id FROM rfc_markdown_index WHERE markdown_id = %s",
            [markdown_id]
        )
        if not existing:
            return json.dumps({"success": False, "error": "markdown_id no encontrado en rfc_markdown_index"})

        if document_name:
            query_insert(
                "UPDATE rfc_markdown_index SET markdown_file_name = %s WHERE markdown_id = %s",
                [document_name, markdown_id]
            )

        log_entry = query_insert_returning(
            """
            INSERT INTO sharepoint_sync_log
                (markdown_id, operation_type, operation_status, target_folder_id,
                 requested_by, requested_at, executed_at, completed_at)
            VALUES (%s, 'create', 'success', %s, %s, now(), now(), now())
            RETURNING sync_log_id
            """,
            [markdown_id, target_folder_id, requested_by or "system"]
        )

        query_insert(
            "UPDATE rfc_markdown_index SET sync_status = 'synced', last_synced_at = now() WHERE markdown_id = %s",
            [markdown_id]
        )

        if requested_by:
            log_audit(
                "sharepoint",
                "DOCUMENT_CREATE",
                requested_by,
                {"markdown_id": markdown_id}
            )

        return json.dumps({
            "success": True,
            "markdown_id": markdown_id,
            "sync_log_id": str(log_entry["sync_log_id"]),
            "status": "created_local",
            "note": "SharePoint real no configurado. Almacenado en indice local PostgreSQL."
        })
    except Exception as e:
        logger.error(f"Error creando documento: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== HERRAMIENTA 5: SharePoint - Editar Documento =====
@mcp.tool()
def tool_sharepoint_update(
    markdown_id: str,
    new_content: str,
    requested_by: Optional[str] = None
) -> str:
    """
    Editar el contenido de un documento en el indice de gestion documental.
    Args: markdown_id (uuid), new_content, requested_by (email)
    """
    try:
        content_bytes = new_content.encode("utf-8")

        rows = query_insert(
            """
            UPDATE rfc_markdown_index
            SET markdown_content = %s,
                markdown_size_bytes = %s,
                markdown_lines = %s,
                last_modified_at = now(),
                last_modified_by = %s,
                sync_status = 'not_synced'
            WHERE markdown_id = %s
            """,
            [content_bytes, len(content_bytes), new_content.count("\n") + 1, requested_by or "system", markdown_id]
        )

        if rows == 0:
            return json.dumps({"success": False, "error": "markdown_id no encontrado"})

        query_insert(
            """
            INSERT INTO sharepoint_sync_log
                (markdown_id, operation_type, operation_status, requested_by, requested_at, executed_at, completed_at)
            VALUES (%s, 'update', 'success', %s, now(), now(), now())
            """,
            [markdown_id, requested_by or "system"]
        )

        if requested_by:
            log_audit(
                "sharepoint",
                "DOCUMENT_UPDATE",
                requested_by,
                {"markdown_id": markdown_id, "content_size": len(new_content)}
            )

        return json.dumps({
            "success": True,
            "markdown_id": markdown_id,
            "status": "updated"
        })
    except Exception as e:
        logger.error(f"Error actualizando documento: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== HERRAMIENTA 6: SharePoint - Mover Documento =====
@mcp.tool()
def tool_sharepoint_move(
    markdown_id: str,
    source_folder_id: Optional[str] = None,
    target_folder_id: Optional[str] = None,
    new_name: Optional[str] = None,
    requested_by: Optional[str] = None
) -> str:
    """
    Registrar el movimiento de un documento a otra carpeta en el indice documental.
    Args: markdown_id (uuid), source_folder_id, target_folder_id, new_name, requested_by (email)
    """
    try:
        existing = query_select(
            "SELECT markdown_id FROM rfc_markdown_index WHERE markdown_id = %s",
            [markdown_id]
        )
        if not existing:
            return json.dumps({"success": False, "error": "markdown_id no encontrado"})

        if new_name:
            query_insert(
                "UPDATE rfc_markdown_index SET markdown_file_name = %s, last_modified_at = now(), last_modified_by = %s WHERE markdown_id = %s",
                [new_name, requested_by or "system", markdown_id]
            )

        query_insert(
            """
            INSERT INTO sharepoint_sync_log
                (markdown_id, operation_type, operation_status, source_path, target_path,
                 target_folder_id, requested_by, requested_at, executed_at, completed_at)
            VALUES (%s, 'move', 'success', %s, %s, %s, %s, now(), now(), now())
            """,
            [markdown_id, source_folder_id, target_folder_id, target_folder_id, requested_by or "system"]
        )

        if requested_by:
            log_audit(
                "sharepoint",
                "DOCUMENT_MOVE",
                requested_by,
                {"markdown_id": markdown_id, "target": target_folder_id}
            )

        return json.dumps({
            "success": True,
            "markdown_id": markdown_id,
            "new_location": target_folder_id or "root"
        })
    except Exception as e:
        logger.error(f"Error moviendo documento: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== HERRAMIENTA 7: SharePoint - Leer Documento =====
@mcp.tool()
def tool_sharepoint_read(
    markdown_id: str,
    requested_by: Optional[str] = None
) -> str:
    """
    Leer el contenido real de un documento del indice de gestion documental.
    Args: markdown_id (uuid), requested_by (email)
    """
    try:
        result = query_select(
            "SELECT markdown_content, markdown_file_name, markdown_size_bytes FROM rfc_markdown_index WHERE markdown_id = %s",
            [markdown_id]
        )

        if not result:
            return json.dumps({"success": False, "error": "markdown_id no encontrado"})

        row = result[0]
        content_bytes = row.get("markdown_content")
        content = bytes(content_bytes).decode("utf-8") if content_bytes else ""

        if requested_by:
            log_audit(
                "sharepoint",
                "DOCUMENT_READ",
                requested_by,
                {"markdown_id": markdown_id}
            )

        return json.dumps({
            "success": True,
            "markdown_id": markdown_id,
            "file_name": row.get("markdown_file_name"),
            "content": content,
            "size": row.get("markdown_size_bytes")
        })
    except Exception as e:
        logger.error(f"Error leyendo documento: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== HERRAMIENTA 8: Sincronizar a SharePoint =====
@mcp.tool()
def tool_sync_to_sharepoint(
    markdown_id: str,
    target_folder_id: Optional[str] = None,
    requested_by: Optional[str] = None
) -> str:
    """
    Sincronizar documento Markdown (SharePoint no conectado aun: sincroniza
    contra el indice local en PostgreSQL).
    Args: markdown_id (uuid), target_folder_id, requested_by (email)
    """
    try:
        existing = query_select(
            "SELECT markdown_id FROM rfc_markdown_index WHERE markdown_id = %s",
            [markdown_id]
        )
        if not existing:
            return json.dumps({"success": False, "error": "markdown_id no encontrado"})

        log_entry = query_insert_returning(
            """
            INSERT INTO sharepoint_sync_log
                (markdown_id, operation_type, operation_status, target_folder_id,
                 requested_by, requested_at, executed_at, completed_at)
            VALUES (%s, 'update', 'success', %s, %s, now(), now(), now())
            RETURNING sync_log_id
            """,
            [markdown_id, target_folder_id, requested_by or "system"]
        )

        query_insert(
            "UPDATE rfc_markdown_index SET sync_status = 'synced', last_synced_at = now() WHERE markdown_id = %s",
            [markdown_id]
        )

        if requested_by:
            log_audit(
                "sharepoint",
                "SYNC_COMPLETED",
                requested_by,
                {"markdown_id": markdown_id, "sync_log_id": str(log_entry["sync_log_id"])}
            )

        return json.dumps({
            "success": True,
            "sync_log_id": str(log_entry["sync_log_id"]),
            "markdown_id": markdown_id,
            "status": "synced_local",
            "note": "SharePoint real no configurado. Sincronizado en indice local PostgreSQL."
        })
    except Exception as e:
        logger.error(f"Error en sincronizacion: {e}")
        return json.dumps({"success": False, "error": str(e)})

# ===== MAIN: INICIAR SERVIDOR =====
def main():
    """Iniciar el MCP Server"""
    try:
        # Validar configuración
        validate_config()

        # Conectar a base de datos
        get_db_pool()

        logger.info(f"✓ {MCP_SERVER_NAME} v{MCP_SERVER_VERSION} iniciando...")
        logger.info(f"✓ 8 herramientas cargadas exitosamente")

        # Ejecutar servidor MCP
        mcp.run()

    except KeyboardInterrupt:
        logger.info("✓ Servidor detenido por el usuario")
    except Exception as e:
        logger.error(f"✗ Error fatal: {e}")
        sys.exit(1)
    finally:
        close_db_pool()
        logger.info("✓ Conexiones cerradas")

if __name__ == "__main__":
    main()
